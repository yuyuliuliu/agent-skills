# -*- coding: utf-8 -*-
"""防误读注释工具：给 Excel 表加说明文字 + 单元格批注 + 列头改名，并处理 Excel 文件锁。

设计原则（见 data-semantic-guard skill）：
1. 改前先备份原文件
2. 不按标签字面理解维度；真实语义需由用户确认后再落笔
3. 加防误读注释：顶部说明文字 + 歧义标签首格批注 + 必要时列头改名
4. Excel 占用导致 save 失败时，原文件未损，生成「-预览版」副本过渡

用法：直接改 __main__ 里的参数后运行；或按需 import annotate / safe_save 复用。
"""
import os
import shutil
import openpyxl
from openpyxl.comments import Comment


def annotate(path, sheet,
             note_cell=None, note_text=None,
             comment_cell=None, comment_text=None,
             rename_cell=None, rename_text=None,
             author="Agent"):
    """在指定 sheet 上加防误读注释，返回 workbook 对象（不保存）。"""
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb[sheet]
    if note_cell and note_text:
        old = ws[note_cell].value or ""
        if note_text not in str(old):
            ws[note_cell] = str(old) + note_text
            print(f"[{note_cell}] 已追加说明")
    if comment_cell and comment_text:
        ws[comment_cell].comment = Comment(comment_text, author)
        print(f"[{comment_cell}] 已加批注")
    if rename_cell and rename_text:
        ws[rename_cell] = rename_text
        print(f"[{rename_cell}] 列头改为 {rename_text!r}")
    return wb


def safe_save(wb, path):
    """保存；遇 Excel 文件锁时生成「-预览版」副本，返回最终写入路径。"""
    try:
        wb.save(path)
        print("保存完成:", path)
        return path
    except PermissionError:
        preview = path[:-5] + "-预览版.xlsx" if path.endswith(".xlsx") else path + "-预览版.xlsx"
        wb.save(preview)
        print(f"[WARN] 原文件被 Excel 占用，已生成预览版：{preview}")
        print("请关闭 Excel 后由我写回原文件并清理预览版。")
        return preview


if __name__ == "__main__":
    import sys
    # ===== 用法：python add_antimisread_annotation.py <你的xlsx路径> =====
    if len(sys.argv) < 2:
        print('用法: python add_antimisread_annotation.py <input.xlsx>')
        print('（下方 note/comment 参数为染厂台账案例，按你的业务自行修改）')
        sys.exit(1)
    FILE = sys.argv[1]

    # 改前先备份
    bak_dir = os.path.join(os.path.dirname(FILE), ".bak_防误读注释")
    os.makedirs(bak_dir, exist_ok=True)
    shutil.copy(FILE, os.path.join(bak_dir, os.path.basename(FILE)))
    print("备份完成:", bak_dir)

    wb = annotate(
        FILE, "日染量矩阵",
        note_cell="A2",
        note_text="；表头『红/蓝/黑』为纸记天数轮次记号（非染色颜色），7月实际全为杂色、无黑色品类。",
        comment_cell="B5",
        comment_text=("红/蓝/黑 = 纸记天数轮次记号，不是染的颜色。\n"
                      "纸本台账每天轮流标一个记号，只为方便数天数、对账用。\n"
                      "请勿把『黑』列误读成『黑色染量』。"),
    )
    safe_save(wb, FILE)
